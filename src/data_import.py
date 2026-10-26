"""Atomic, validated CSV/XLSX imports into the persistent data model.

Sales history is append-only (identical reimports are harmless). Master data and
inventory are upserts, with before/after audit evidence for every changed row.
"""
from __future__ import annotations

import csv
import io
import json
import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from src.database import Connection, now_utc
from src.validation import SKU_PATTERN

DATASET_HEADERS: dict[str, list[str]] = {
    'sales': ['sale_date', 'sku', 'units'],
    'inventory': ['sku', 'on_hand', 'on_order', 'reserved_qty'],
    'products': ['sku', 'product_name', 'supplier', 'lead_time_days',
                 'safety_stock', 'pack_size', 'unit_cost_vnd', 'on_hand',
                 'on_order', 'reserved_qty'],
    'suppliers': ['name', 'lead_time_days', 'contact_email'],
}
_REQUIRED = {
    'sales': {'sale_date', 'sku', 'units'},
    'inventory': {'sku', 'on_hand', 'on_order'},
    'products': {'sku', 'product_name', 'supplier', 'lead_time_days',
                 'safety_stock', 'pack_size', 'unit_cost_vnd'},
    'suppliers': {'name'},
}
_NUMERIC = {'units', 'on_hand', 'on_order', 'reserved_qty', 'lead_time_days',
            'safety_stock', 'pack_size', 'unit_cost_vnd'}
MAX_FILE_BYTES = 2_000_000
MAX_ROWS = 10_000


def _missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip()) or (
        isinstance(value, float) and math.isnan(value))


def _read_rows(raw: bytes, filename: str, dataset: str) -> list[dict[str, Any]]:
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise ValueError('File must contain data and be no larger than 2 MB')
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == '.csv':
            try:
                contents = raw.decode('utf-8-sig')
            except UnicodeDecodeError as exc:
                raise ValueError('CSV must be UTF-8 encoded') from exc
            reader = csv.reader(io.StringIO(contents), strict=True)
            headers = next(reader, [])
            rows = []
            for row in reader:
                if not row:
                    continue
                if len(row) != len(headers):
                    raise ValueError('Every row must have the same number of columns as the header')
                rows.append(row)
                if len(rows) > MAX_ROWS:
                    raise ValueError('File row limit: 10000')
        elif suffix == '.xlsx':
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                if len(workbook.worksheets) != 1:
                    raise ValueError('Excel imports must contain exactly one worksheet')
                worksheet = workbook.worksheets[0]
                iterator = worksheet.iter_rows(values_only=True)
                headers = list(next(iterator, []))
                rows = []
                for values in iterator:
                    if all(_missing(value) for value in values):
                        continue
                    rows.append(list(values))
                    if len(rows) > MAX_ROWS:
                        raise ValueError('File row limit: 10000')
            finally:
                workbook.close()
        else:
            raise ValueError('Supported formats are .csv and .xlsx')
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('File could not be read as a valid CSV/XLSX workbook') from exc
    if not headers or any(not isinstance(h, str) for h in headers):
        raise ValueError('File must have a text header row')
    if len(set(headers)) != len(headers):
        raise ValueError('Duplicate column names are not allowed')
    if not _REQUIRED[dataset].issubset(headers) or not set(headers).issubset(DATASET_HEADERS[dataset]):
        raise ValueError(f'Invalid {dataset} schema. Required columns: {", ".join(sorted(_REQUIRED[dataset]))}')
    if not rows:
        raise ValueError('File must contain at least one data row')
    # Object dtype preserves identifiers (e.g. leading-zero SKUs) and avoids
    # coercing blank fields to NaN or turning integer CSV values into floats.
    return pd.DataFrame(rows, columns=headers, dtype=object).to_dict(orient='records')


def _integer(value: Any, field: str, row_number: int) -> int:
    if _missing(value) or isinstance(value, bool):
        raise ValueError(f'Missing or invalid {field} at row {row_number}')
    try:
        parsed = int(value)
    except (ValueError, TypeError, OverflowError) as exc:
        raise ValueError(f'{field} must be an integer at row {row_number}') from exc
    if isinstance(value, float) and (not math.isfinite(value) or value != parsed):
        raise ValueError(f'{field} must be an integer at row {row_number}')
    minimum = 1 if field in {'lead_time_days', 'pack_size'} else 0
    if parsed < minimum or parsed > 1_000_000_000:
        raise ValueError(f'{field} must be between {minimum} and 1000000000 at row {row_number}')
    return parsed


def _validate(rows: list[dict[str, Any]], dataset: str, as_of: date) -> list[dict[str, Any]]:
    seen: set[Any] = set()
    result = []
    for number, row in enumerate(rows, start=2):
        parsed: dict[str, Any] = {}
        for field, value in row.items():
            if field in _NUMERIC:
                if dataset == 'suppliers' and _missing(value):
                    parsed[field] = None
                else:
                    parsed[field] = _integer(value, field, number)
            elif field == 'sale_date':
                try:
                    if isinstance(value, datetime):
                        day = value.date()
                    elif isinstance(value, date):
                        day = value
                    else:
                        value = str(value).strip()
                        day = date.fromisoformat(value)
                        if day.isoformat() != value:
                            raise ValueError('Use YYYY-MM-DD')
                except (ValueError, TypeError, AttributeError) as exc:
                    raise ValueError(f'Invalid date at row {number}; use YYYY-MM-DD') from exc
                if day > as_of:
                    raise ValueError(f'Future date at row {number}')
                parsed[field] = day.isoformat()
            else:
                if _missing(value):
                    if field == 'contact_email':
                        parsed[field] = None
                        continue
                    raise ValueError(f'Missing {field} at row {number}')
                if not isinstance(value, str):
                    raise ValueError(f'{field} must be text at row {number}')
                text = value.strip()
                if field == 'sku' and re.fullmatch(SKU_PATTERN, text) is None:
                    raise ValueError(f'Invalid SKU at row {number}; use 1-64 letters, digits, dot, underscore or hyphen')
                if len(text) > 200 or any(ord(c) < 32 for c in text):
                    raise ValueError(f'Invalid or overly long {field} at row {number}')
                if field == 'contact_email' and ('@' not in text or ' ' in text):
                    raise ValueError(f'Invalid contact_email at row {number}')
                parsed[field] = text
        key = (parsed['sale_date'], parsed['sku']) if dataset == 'sales' else parsed[
            'name' if dataset == 'suppliers' else 'sku']
        if key in seen:
            raise ValueError(f'Duplicate record in file at row {number}')
        seen.add(key)
        result.append(parsed)
    return result


def record_inventory_snapshot(conn: Connection, sku: str, as_of: date,
                              source: str) -> None:
    """Save the actual simulated stock state, never infer it from sales."""
    row = conn.execute('SELECT on_hand,on_order,reserved_qty FROM inventory WHERE sku=?', (sku,)).fetchone()
    if row is None:
        return
    conn.execute('''INSERT INTO inventory_snapshots
        (snapshot_date,sku,on_hand,on_order,reserved_qty,source,created_at)
        VALUES(?,?,?,?,?,?,?) ON CONFLICT(snapshot_date,sku,source) DO UPDATE SET
        on_hand=excluded.on_hand,on_order=excluded.on_order,
        reserved_qty=excluded.reserved_qty,created_at=excluded.created_at''',
        (as_of.isoformat(), sku, row['on_hand'], row['on_order'], row['reserved_qty'], source, now_utc()))


def import_dataset(conn: Connection, raw: bytes, filename: str, dataset: str,
                   actor: str = 'demo_operator', as_of: date | None = None) -> int:
    """Validate the whole file, then atomically apply changes and audit evidence.

    Required headers are listed in ``_REQUIRED``; optional headers appear in
    ``DATASET_HEADERS``. Inventory reserved_qty defaults to zero when omitted;
    omitted product stock fields preserve existing values (zero for new SKUs).
    An approved request locks its SKU's on_order against data import changes.
    """
    if dataset not in DATASET_HEADERS:
        raise ValueError('Dataset must be sales, inventory, products, or suppliers')
    if not actor or not actor.strip() or len(actor.strip()) > 100:
        raise ValueError('A valid import actor is required (1-100 characters)')
    as_of = as_of or date.today()
    rows = _validate(_read_rows(raw, filename, dataset), dataset, as_of)
    before_records: list[dict[str, Any]] = []
    after_records: list[dict[str, Any]] = []
    with conn:
        # Serialize against approval and other imports before validating stored
        # state. SQLite otherwise begins a deferred transaction only on write.
        if not conn.in_transaction:
            conn.execute('BEGIN IMMEDIATE')
        skus = {row['sku'] for row in conn.execute('SELECT sku FROM products')}
        suppliers = {row['name'] for row in conn.execute('SELECT name FROM suppliers')}
        for row in rows:
            if dataset in {'sales', 'inventory'} and row['sku'] not in skus:
                raise ValueError(f'Unknown SKU: {row["sku"]}')
            if dataset == 'products' and row['supplier'] not in suppliers:
                raise ValueError(f'Unknown supplier: {row["supplier"]}; import suppliers first')
            if dataset == 'sales':
                stored = conn.execute('SELECT * FROM sales WHERE sale_date=? AND sku=?',
                                      (row['sale_date'], row['sku'])).fetchone()
                if stored is not None:
                    if stored['units'] != row['units']:
                        raise ValueError(f'Conflicting sales already stored for {row["sale_date"]}/{row["sku"]}')
                    continue
                values = row
            else:
                key = 'name' if dataset == 'suppliers' else 'sku'
                stored = conn.execute(f'SELECT * FROM {dataset} WHERE {key}=?', (row[key],)).fetchone()
                if dataset == 'products':
                    values = dict(stored) if stored else {'on_hand': 0, 'on_order': 0, 'reserved_qty': 0}
                    values.update(row)
                elif dataset == 'inventory':
                    values = {**row, 'reserved_qty': row.get('reserved_qty', 0)}
                else:
                    values = dict(stored) if stored else {'lead_time_days': None, 'contact_email': None}
                    values.update(row)
                if stored and all(stored[field] == value for field, value in values.items()):
                    continue
                if dataset in {'inventory', 'products'} and stored:
                    approved = conn.execute("SELECT 1 FROM proposals WHERE sku=? AND status='APPROVED' LIMIT 1",
                                            (row['sku'],)).fetchone()
                    if approved and values['on_order'] != stored['on_order']:
                        raise ValueError(f'Cannot change on_order for {row["sku"]} while an APPROVED request is outstanding')
            before_records.append(dict(stored) if stored is not None else {})
            if dataset == 'inventory':
                values['updated_at'] = now_utc()
            fields = list(values)
            placeholders = ','.join('?' for _ in fields)
            if dataset == 'sales':
                conn.execute(f'INSERT INTO sales ({",".join(fields)}) VALUES ({placeholders})',
                             tuple(values[field] for field in fields))
            else:
                key = 'name' if dataset == 'suppliers' else 'sku'
                updates = ','.join(f'{field}=excluded.{field}' for field in fields if field != key)
                conn.execute(f'INSERT INTO {dataset} ({",".join(fields)}) VALUES ({placeholders}) '
                             f'ON CONFLICT({key}) DO UPDATE SET {updates}', tuple(values[field] for field in fields))
            after_records.append(values)
            if dataset in {'inventory', 'products'}:
                record_inventory_snapshot(conn, row['sku'], as_of, f'import:{dataset}')
        if after_records:
            conn.execute('''INSERT INTO audit_logs(action,actor,comment,before_json,after_json,created_at)
                VALUES(?,?,?,?,?,?)''', ('DATA_IMPORTED', actor.strip(),
                f'{dataset}: {len(after_records)} changed row(s) from {Path(filename).name}',
                json.dumps({'dataset': dataset, 'records': before_records}, ensure_ascii=False),
                json.dumps({'dataset': dataset, 'records': after_records}, ensure_ascii=False), now_utc()))
    return len(after_records)
