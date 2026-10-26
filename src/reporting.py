"""Portable, formula-safe CSV and Excel reports for Excel and Power BI."""
from __future__ import annotations

import csv
import io
from datetime import date
from typing import Any

import pandas as pd


def safe_spreadsheet_cell(value: Any) -> Any:
    """Escape formula-looking text to reduce CSV formula injection risk."""
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def make_csv(rows: list[dict[str, Any]], fields: list[str]) -> bytes:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction='ignore')
    writer.writeheader()
    for row in rows:
        writer.writerow({k: safe_spreadsheet_cell(row.get(k, '')) for k in fields})
    return ('\ufeff' + output.getvalue()).encode('utf-8')


def make_excel(sheets: dict[str, list[dict[str, Any]]]) -> bytes:
    """Return a workbook with one table per sheet and text formulas escaped."""
    from openpyxl.utils import get_column_letter

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        for name, rows in (sheets or {'report': []}).items():
            records = [{key: safe_spreadsheet_cell(value) for key, value in row.items()}
                       for row in rows]
            frame = pd.DataFrame(records) if records else pd.DataFrame(columns=report_fields(name))
            frame.to_excel(writer, sheet_name=name[:31], index=False)
            worksheet = writer.sheets[name[:31]]
            worksheet.freeze_panes = 'A2'
            if len(frame.columns):
                worksheet.auto_filter.ref = worksheet.dimensions
                for index, column in enumerate(frame.columns, start=1):
                    worksheet.column_dimensions[get_column_letter(index)].width = min(
                        48, max(16, len(str(column)) + 3)
                    )
    return output.getvalue()


REPORT_TABLES = (
    'suppliers', 'products', 'inventory', 'sales', 'proposals',
    'approval_history', 'audit_logs', 'exceptions', 'inventory_snapshots',
)


def report_fields(dataset: str) -> list[str]:
    """Retain discoverable headers even when a normalized report table is empty."""
    from src.models import Base

    if dataset in REPORT_TABLES:
        return [column.name for column in Base.metadata.tables[dataset].columns]
    if dataset == 'replenishment':
        return ['assessment_date', 'sku', 'product_name', 'supplier', 'risk', 'on_hand', 'on_order',
                'reserved_qty', 'safety_stock', 'lead_time_days', 'avg_daily_demand',
                'inventory_position', 'reorder_point', 'target_stock', 'quantity',
                'estimated_cost_vnd', 'demand_change_pct']
    return []


def build_report(conn: Any, as_of: date | None = None) -> dict[str, list[dict[str, Any]]]:
    """Read normalized tables and the dated assessment without changing records."""
    from src.procurement import get_replenishment_table

    sheets = {
        table: [dict(row) for row in conn.execute(f'SELECT * FROM {table}').fetchall()]
        for table in REPORT_TABLES
    }
    from src.procurement import business_today
    assessment_date = (as_of or business_today()).isoformat()
    sheets['replenishment'] = [
        {'assessment_date': assessment_date, **row}
        for row in get_replenishment_table(conn, as_of)
    ]
    return sheets
