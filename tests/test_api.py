"""HTTP integration tests use temporary databases and local agent behavior."""
import csv
import io
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src import local_agent
from src.api import create_app


AS_OF = '2026-10-26'
AGENT_STATUS = {
    'provider': 'local_rules', 'mode': 'deterministic', 'available': True,
    'model': None, 'endpoint': None, 'message': 'Local advisory mode',
}


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(local_agent, 'agent_status', lambda: dict(AGENT_STATUS))
    app = create_app(tmp_path / 'workspace.db', tmp_path / 'missing-build', date(2026, 10, 26))
    with TestClient(app, raise_server_exceptions=False) as test_client:
        yield test_client


def workspace(client):
    response = client.get('/api/workspace', params={'as_of': AS_OF})
    assert response.status_code == 200, response.text
    return response.json()


def test_workspace_and_health_expose_persisted_synthetic_data(client):
    assert client.get('/api/health').json()['status'] == 'ok'
    data = workspace(client)
    assert data['as_of'] == AS_OF
    assert len(data['inventory']) == 6
    assert len(data['sales']) == len(data['snapshots']) == 504
    assert len(data['suppliers']) == 3
    assert data['agent'] == AGENT_STATUS
    assert data['proposals'] == data['approvals'] == []
    assert data['settings']['operator_name']


def test_generate_decide_update_and_reject_are_real_workflow(client):
    response = client.patch('/api/inventory/CF-001', json={
        'on_hand': 12, 'on_order': 0, 'reserved_qty': 5, 'actor': 'API tester', 'as_of': AS_OF,
    })
    assert response.status_code == 200
    assert response.json()['imported'] == 1
    payload = {'actor': 'API tester', 'as_of': AS_OF}
    assert client.post('/api/requests/generate', json=payload).json()['created'] >= 2
    assert client.post('/api/requests/generate', json=payload).json()['created'] == 0
    before = workspace(client)
    approved = next(row for row in before['proposals'] if row['sku'] == 'CF-001')
    rejected = next(row for row in before['proposals'] if row['sku'] != 'CF-001')
    decision = {'decision': 'APPROVED', 'actor': 'API reviewer', 'comment': 'Checked synthetic stock and delivery assumptions'}
    assert client.post(f'/api/requests/{approved["id"]}/decision', json=decision).status_code == 200
    assert client.post(f'/api/requests/{approved["id"]}/decision', json=decision).status_code == 400
    decision.update(decision='REJECTED', comment='Defer material while confirming demand')
    assert client.post(f'/api/requests/{rejected["id"]}/decision', json=decision).status_code == 200
    after = workspace(client)
    stock = next(row for row in after['inventory'] if row['sku'] == 'CF-001')
    assert stock['on_order'] == approved['quantity']
    assert len(after['approvals']) == 2
    assert {'APPROVED', 'REJECTED'} <= {row['action'] for row in after['audit']}
    locked = client.patch('/api/inventory/CF-001', json={
        'on_hand': 12, 'on_order': 0, 'reserved_qty': 5, 'actor': 'API tester', 'as_of': AS_OF,
    })
    assert locked.status_code == 400
    assert 'APPROVED' in locked.json()['detail']


def test_upload_csv_xlsx_and_templates(client):
    template = client.get('/api/templates/inventory.csv')
    assert template.status_code == 200
    assert template.content.decode('utf-8-sig').strip() == 'sku,on_hand,on_order,reserved_qty'
    raw = b'sale_date,sku,units\n2026-07-01,CF-001,10\n'
    response = client.post('/api/import/sales', params={'as_of': AS_OF},
                           files={'file': ('sales.csv', raw, 'text/csv')}, data={'actor': 'Uploader'})
    assert response.status_code == 200
    assert response.json()['imported'] == 1
    from src.reporting import make_excel
    workbook = make_excel({'inventory': [{'sku': 'CF-002', 'on_hand': 80, 'on_order': 0, 'reserved_qty': 2}]})
    response = client.post('/api/import/inventory', params={'as_of': AS_OF},
                           files={'file': ('inventory.xlsx', workbook)}, data={'actor': 'Uploader'})
    assert response.status_code == 200
    assert response.json()['imported'] == 1
    broken = client.post('/api/import/inventory', files={
        'file': ('broken.csv', b'sku,on_hand,on_order\nCF-001,1,0,extra\n'),
    })
    assert broken.status_code == 400
    oversized = client.post('/api/import/sales', files={'file': ('large.csv', b'x' * 2_000_001)})
    assert oversized.status_code == 413


def test_csv_and_excel_reports_contain_actual_snapshot_values(client):
    assert client.post('/api/requests/generate', json={'actor': 'Tester', 'as_of': AS_OF}).status_code == 200
    response = client.get('/api/reports/requests.csv', params={'as_of': AS_OF})
    assert response.status_code == 200
    assert 'attachment;' in response.headers['content-disposition']
    records = list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))
    assert len(records) >= 2
    assert all(int(row['estimated_cost_vnd']) == int(row['quantity']) * int(row['unit_cost_vnd']) for row in records)
    assert all(row['supplier'] and row['product_name'] for row in records)
    response = client.get('/api/reports/workbook.xlsx', params={'as_of': AS_OF})
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    try:
        assert {'inventory', 'sales', 'proposals', 'audit_logs', 'replenishment'} <= set(workbook.sheetnames)
        assert workbook['sales'].max_row == 505
    finally:
        workbook.close()
    assert client.get('/api/reports/unknown.csv').status_code == 404
    assert client.get('/api/templates/unknown.csv').status_code == 404


def test_settings_survive_service_restart_and_changes_are_audited(tmp_path, monkeypatch):
    monkeypatch.setattr(local_agent, 'agent_status', lambda: dict(AGENT_STATUS))
    path = tmp_path / 'persistent.db'
    app = create_app(path, seed_date=date(2026, 10, 26))
    with TestClient(app) as first:
        response = first.patch('/api/settings', json={'operator_name': 'Portfolio Reviewer', 'workspace_name': 'Synthetic Operations'})
        assert response.status_code == 200
        assert response.json()['workspace_name'] == 'Synthetic Operations'
    with TestClient(create_app(path)) as second:
        assert second.get('/api/settings').json() == {
            'operator_name': 'Portfolio Reviewer', 'workspace_name': 'Synthetic Operations',
        }
        assert any(row['action'] == 'SETTINGS_UPDATED' for row in workspace(second)['audit'])


@pytest.mark.parametrize('method,path,payload', [
    ('post', '/api/requests/generate', {'actor': '', 'as_of': AS_OF}),
    ('post', '/api/requests/generate', {'actor': 'Tester', 'as_of': '20261026'}),
    ('post', '/api/requests/1/decision', {'decision': 'APPROVED', 'actor': 'Tester', 'comment': ''}),
    ('patch', '/api/inventory/CF-001', {'on_hand': -1, 'on_order': 0, 'reserved_qty': 0, 'actor': 'Tester'}),
    ('patch', '/api/inventory/CF-001', {'on_hand': 1.5, 'on_order': 0, 'reserved_qty': 0, 'actor': 'Tester'}),
    ('patch', '/api/settings', {'operator_name': ' '}),
    ('patch', '/api/settings', {'workspace_name': None}),
    ('patch', '/api/settings', {'unexpected': 'value'}),
    ('post', '/api/agent/chat', {'message': ' '}),
])
def test_invalid_inputs_are_clear_json_errors(client, method, path, payload):
    response = getattr(client, method)(path, json=payload)
    assert response.status_code == 422
    assert 'detail' in response.json()


def test_local_origin_checks_block_external_browser_mutations(client):
    response = client.patch('/api/settings', json={'workspace_name': 'Unwanted change'},
                            headers={'Origin': 'https://external.example'})
    assert response.status_code == 403
    assert client.get('/api/settings').json()['workspace_name'] != 'Unwanted change'
    allowed = client.options('/api/settings', headers={
        'Origin': 'http://localhost:5173', 'Access-Control-Request-Method': 'PATCH',
    })
    assert allowed.status_code == 200
    assert allowed.headers['access-control-allow-origin'] == 'http://localhost:5173'


def test_unknown_resources_and_missing_frontend_do_not_return_spa_html(client):
    assert client.get('/').status_code == 503
    for path in ('/api/unknown', '/api', '/assets/missing.js', '/api/templates/not-a-dataset.csv'):
        response = client.get(path)
        assert response.status_code == 404
        assert response.headers['content-type'].startswith('application/json')
    assert client.get('/api/workspace', params={'as_of': '20261026'}).status_code == 422
    assert client.patch('/api/inventory/UNKNOWN', json={
        'on_hand': 1, 'on_order': 0, 'reserved_qty': 0, 'actor': 'Tester',
    }).status_code == 404


def test_frontend_assets_and_spa_routes_are_served(tmp_path):
    distribution = tmp_path / 'dist'
    distribution.mkdir()
    (distribution / 'assets').mkdir()
    (distribution / 'index.html').write_text('<html><body>ProcureFlow test shell</body></html>', encoding='utf-8')
    (distribution / 'assets' / 'app.js').write_text('console.log("loaded")', encoding='utf-8')
    with TestClient(create_app(tmp_path / 'static.db', distribution)) as client:
        assert 'ProcureFlow test shell' in client.get('/').text
        assert 'ProcureFlow test shell' in client.get('/inventory').text
        assert client.get('/assets/app.js').status_code == 200
        assert client.get('/api/does-not-exist').status_code == 404


def test_unexpected_errors_do_not_expose_credentials(client, monkeypatch):
    import src.api as api

    def fail(*args, **kwargs):
        raise RuntimeError('SECRET_CREDENTIAL_MUST_NOT_APPEAR')

    monkeypatch.setattr(api, 'create_drafts', fail)
    response = client.post('/api/requests/generate', json={'actor': 'Tester', 'as_of': AS_OF})
    assert response.status_code == 500
    assert 'SECRET_CREDENTIAL' not in response.text
    assert 'detail' in response.json()


def test_chat_route_delegates_and_conversation_routes_load_saved_history(client, monkeypatch):
    """Agent implementation has its own real-provider tests; verify API wiring."""
    calls = []

    def chat(conn, message, conversation_id=None, as_of=None, sku=None):
        calls.append((message, conversation_id, as_of, sku))
        return {'conversation_id': 'conversation-test', 'message': 'Review stock and pending deliveries.',
                'source': 'LOCAL_RULES', 'actions': []}

    monkeypatch.setattr(local_agent, 'chat', chat)
    monkeypatch.setattr(local_agent, 'list_conversations', lambda conn: [
        {'id': 'conversation-test', 'title': 'Inventory review', 'created_at': '2026-10-26', 'updated_at': '2026-10-26'},
    ])
    monkeypatch.setattr(local_agent, 'get_messages', lambda conn, conversation_id: [
        {'id': 1, 'role': 'assistant', 'content': 'Review stock.', 'metadata': {}, 'created_at': '2026-10-26'},
    ])
    response = client.post('/api/agent/chat', json={'message': 'Review CF-001', 'sku': 'CF-001', 'as_of': AS_OF})
    assert response.status_code == 200
    assert response.json()['conversation_id'] == 'conversation-test'
    assert calls == [('Review CF-001', None, date(2026, 10, 26), 'CF-001')]
    assert client.get('/api/agent/status').json() == AGENT_STATUS
    assert client.get('/api/agent/conversations').json()[0]['id'] == 'conversation-test'
    assert client.get('/api/agent/conversations/conversation-test/messages').json()[0]['role'] == 'assistant'


def test_real_offline_agent_persists_chat_without_changing_business_data(client, monkeypatch):
    monkeypatch.setenv('LOCAL_LLM_ENABLED', 'false')
    monkeypatch.setattr(local_agent, 'business_today', lambda: date(2026, 10, 26))
    before = workspace(client)
    response = client.post('/api/agent/chat', json={
        'message': 'Why reorder CF-001?', 'sku': 'CF-001', 'as_of': AS_OF,
    })
    assert response.status_code == 200, response.text
    answer = response.json()
    assert answer['mode'] == 'local_tools'
    assert 'CF-001' in answer['answer']
    assert answer['tools'] and answer['sources']
    conversation_id = answer['conversation_id']
    conversations = client.get('/api/agent/conversations').json()
    assert conversations[0]['id'] == conversation_id
    history = client.get(f'/api/agent/conversations/{conversation_id}/messages').json()
    assert [row['role'] for row in history] == ['user', 'assistant']
    # An omitted date follows the conversation's retained analysis context.
    monkeypatch.setattr(local_agent, 'business_today', lambda: date(2026, 10, 27))
    followup = client.post('/api/agent/chat', json={
        'message': 'What is its budget?', 'conversation_id': conversation_id,
    })
    assert followup.status_code == 200
    assert '2026-10-26' in followup.json()['answer']
    after = workspace(client)
    for name in ('inventory', 'proposals', 'audit', 'approvals', 'sales', 'snapshots'):
        assert after[name] == before[name]
    invalid = client.post('/api/agent/chat', json={'message': 'Review inventory', 'sku': 'UNKNOWN'})
    assert invalid.status_code == 400
