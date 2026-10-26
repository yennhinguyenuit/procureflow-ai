"""The offline portfolio scenario persists actual inputs, decisions, and exports."""
import csv
import io
import json
from datetime import date

from openpyxl import load_workbook

from scripts.demo import run_demo
from scripts.export_samples import export_samples
from src.data_import import import_dataset
from src.database import connect, init_db


def test_demo_persists_workflow_and_preserves_previous_run(tmp_path, monkeypatch):
    # A configured API key must still have no effect on this offline scenario.
    monkeypatch.setenv('GEMINI_API_KEY', 'must-never-be-used')
    import src.ai_assistant as assistant

    def no_sdk():
        raise AssertionError('Offline demo attempted to create an AI client')

    monkeypatch.setattr(assistant, 'summarize', no_sdk)
    first = run_demo(tmp_path, date(2026, 10, 26))
    original_db = (first / 'procureflow.db').read_bytes()
    second = run_demo(tmp_path, date(2026, 10, 26))
    assert first != second
    assert (first / 'procureflow.db').read_bytes() == original_db
    manifest = json.loads((first / 'run_manifest.json').read_text(encoding='utf-8'))
    assert manifest['persisted_verification'] is True
    assert manifest['imported_changed_rows'] == {
        'suppliers': 3, 'products': 6, 'inventory': 6, 'sales': 504, 'shortage': 1,
    }
    assert manifest['drafts_on_repeated_generation'] == 0
    assert manifest['identical_sales_import_changed_rows'] == 0
    assert manifest['on_order_after_approval'] - manifest['on_order_before_approval'] == manifest['approved_quantity']
    conn = connect(first / 'procureflow.db')
    try:
        statuses = {row['status']: row['count'] for row in conn.execute(
            'SELECT status,COUNT(*) AS count FROM proposals GROUP BY status')}
        assert statuses['APPROVED'] == statuses['REJECTED'] == 1
        assert conn.execute('SELECT COUNT(*) FROM approval_history').fetchone()[0] == 2
        assert conn.execute("SELECT COUNT(*) FROM audit_logs WHERE action='DATA_IMPORTED'").fetchone()[0] == 5
    finally:
        conn.close()
    workbook = load_workbook(first / 'reports.xlsx', read_only=True)
    try:
        assert set(workbook.sheetnames) == set(manifest['report_rows'])
        assert workbook['sales'].max_row == 505
        assert workbook['approval_history'].max_row == 3
    finally:
        workbook.close()
    assert 'human approval' in (first / 'offline_insights.txt').read_text(encoding='utf-8')


def test_exported_csv_and_excel_samples_are_importable(tmp_path):
    paths = export_samples(tmp_path, date(2026, 10, 26))
    conn = connect(':memory:')
    try:
        init_db(conn)
        for dataset, expected in (('suppliers', 3), ('products', 6), ('inventory', 6), ('sales', 504)):
            path = paths[dataset]
            assert import_dataset(conn, path.read_bytes(), path.name, dataset,
                                  as_of=date(2026, 10, 26)) == expected
        for dataset in ('sales', 'inventory'):
            path = paths[f'{dataset}_xlsx']
            assert import_dataset(conn, path.read_bytes(), path.name, dataset,
                                  as_of=date(2026, 10, 26)) == 0
        assert import_dataset(conn, paths['demo_sales'].read_bytes(), 'demo_sales.csv', 'sales',
                              as_of=date(2026, 10, 26)) == 0
    finally:
        conn.close()
    rows = list(csv.DictReader(io.StringIO(paths['inventory'].read_text(encoding='utf-8-sig'))))
    assert set(rows[0]) == {'sku', 'on_hand', 'on_order', 'reserved_qty'}
