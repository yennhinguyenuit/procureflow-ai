"""Exercise the rendered application and exported artifacts with temporary data."""
from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

from src.database import connect
from src.reporting import build_report, make_excel


@pytest.fixture
def app(tmp_path, monkeypatch):
    database = tmp_path / 'ui-test.db'
    monkeypatch.setenv('PROCUREFLOW_DB_PATH', str(database))
    # Explicit empty configuration keeps the test offline even with a local .env.
    monkeypatch.setenv('GEMINI_API_KEY', '')
    application = AppTest.from_file(str(Path(__file__).parents[1] / 'app.py'), default_timeout=30)
    application.run()
    _assert_rendered(application)
    return application, database


def _assert_rendered(app):
    assert not app.exception, [error.value for error in app.exception]


def _navigate(app, page):
    app.sidebar.radio[0].set_value(page).run()
    _assert_rendered(app)
    assert app.title[0].value == page


def _button(app, label):
    return next(button for button in app.button if button.label == label)


def test_all_pages_filters_and_report_controls(app):
    application, database = app
    assert int(application.metric[0].value) >= 6
    _navigate(application, 'Inventory')
    application.sidebar.multiselect(key='sku_filter').set_value(['CF-001']).run()
    _assert_rendered(application)
    assert set(application.dataframe[0].value['sku']) == {'CF-001'}
    application.sidebar.multiselect(key='sku_filter').set_value([]).run()
    for page in ('Procurement Requests', 'AI Insights', 'Analytics', 'Audit Logs', 'Settings'):
        _navigate(application, page)
    application.text_input(key='operator_input').set_value('portfolio_reviewer').run()
    _navigate(application, 'Procurement Requests')
    _button(application, 'Generate draft purchase requests').click().run()
    _assert_rendered(application)
    conn = connect(database)
    try:
        assert conn.execute("SELECT COUNT(*) FROM audit_logs WHERE actor='portfolio_reviewer'").fetchone()[0] > 0
    finally:
        conn.close()
    _navigate(application, 'Analytics')
    downloads = {element.proto.label for element in application.get('download_button')}
    assert {'Download full Excel report', 'Download selected CSV report'} <= downloads
    _navigate(application, 'Overview')


def test_ui_generate_deduplicate_approve_reject_and_persist(app):
    application, database = app
    _navigate(application, 'Procurement Requests')
    _button(application, 'Generate draft purchase requests').click().run()
    _assert_rendered(application)
    conn = connect(database)
    try:
        count = conn.execute('SELECT COUNT(*) FROM proposals').fetchone()[0]
        assert count >= 2
    finally:
        conn.close()
    _button(application, 'Generate draft purchase requests').click().run()
    _assert_rendered(application)
    conn = connect(database)
    try:
        assert conn.execute('SELECT COUNT(*) FROM proposals').fetchone()[0] == count
    finally:
        conn.close()

    # A missing reason is a visible validation error and cannot change a draft.
    _button(application, 'Approve selected draft').click().run()
    assert application.error
    _assert_rendered(application)
    application.text_area(key='review_reason').set_value('Reviewed demand and supplier coverage.').run()
    _button(application, 'Approve selected draft').click().run()
    _assert_rendered(application)
    application.text_area(key='review_reason').set_value('Rejected after checking the demo plan.').run()
    _button(application, 'Reject selected draft').click().run()
    _assert_rendered(application)
    conn = connect(database)
    try:
        statuses = {r['status'] for r in conn.execute('SELECT status FROM proposals')}
        assert {'APPROVED', 'REJECTED'} <= statuses
        history = conn.execute('SELECT decision FROM approval_history').fetchall()
        assert {r['decision'] for r in history} == {'APPROVED', 'REJECTED'}
        approved = conn.execute("SELECT sku,quantity FROM proposals WHERE status='APPROVED'").fetchone()
        assert conn.execute('SELECT on_order FROM inventory WHERE sku=?', (approved['sku'],)).fetchone()[0] >= approved['quantity']
        workbook = openpyxl.load_workbook(io.BytesIO(make_excel(build_report(conn))))
        assert {'products', 'suppliers', 'inventory', 'sales', 'proposals', 'approval_history',
                'audit_logs', 'exceptions', 'inventory_snapshots', 'replenishment'} <= set(workbook.sheetnames)
        assert workbook['approval_history'].max_row == 3
    finally:
        conn.close()
    # A separate browser session reads the same committed decisions.
    second_session = AppTest.from_file(str(Path(__file__).parents[1] / 'app.py'), default_timeout=30).run()
    _navigate(second_session, 'Audit Logs')
    rendered_history = next(frame.value for frame in second_session.dataframe if 'decision' in frame.value.columns)
    assert set(rendered_history['decision']) == {'APPROVED', 'REJECTED'}


def test_ui_offline_ai_summary_and_stale_filter_warning(app):
    application, _ = app
    _navigate(application, 'AI Insights')
    _button(application, 'Generate inventory risk summary').click().run()
    _assert_rendered(application)
    assert any('RULE_BASED_FALLBACK' in item.value for item in application.caption)
    assert any('human approval' in item.value.lower() for item in application.markdown)
    application.sidebar.multiselect(key='sku_filter').set_value(['CF-001']).run()
    _assert_rendered(application)
    assert any('changed since this analysis' in item.value for item in application.warning)


def test_excel_export_preserves_numeric_values_and_escapes_formulas():
    raw = make_excel({'suppliers': [{'name': '=HYPERLINK("https://example.invalid")', 'lead_time_days': 5}],
                      'exceptions': []})
    book = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    sheet = book['suppliers']
    assert sheet['A2'].value.startswith("'=HYPERLINK")
    assert sheet['A2'].data_type == 's'
    assert sheet['B2'].value == 5
    assert sheet['B2'].data_type == 'n'
    assert 'exceptions' in book.sheetnames
