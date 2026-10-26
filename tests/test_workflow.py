"""Workflow invariants, independent connections, and persistent import-to-report flow."""
from __future__ import annotations

import io
import json
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from datetime import date

import openpyxl
import pytest

from src.data_import import import_dataset
from src.database import connect, init_db
from src.procurement import (
    calculate_replenishment, create_drafts, decide, get_approval_history,
    get_audit_logs, get_exceptions, get_proposals, get_replenishment_table,
)
from src.reporting import build_report, make_csv, make_excel
from src.seed import seed_if_empty

AS_OF = date(2026, 10, 26)


@pytest.fixture
def db(tmp_path):
    path = tmp_path / 'workflow.db'
    conn = connect(path)
    init_db(conn)
    seed_if_empty(conn, AS_OF)
    yield conn, path
    conn.close()


def test_reserved_stock_and_reorder_boundary():
    product = {'sku': 'X', 'on_hand': 75, 'on_order': 5, 'reserved_qty': 10,
               'lead_time_days': 5, 'safety_stock': 20, 'pack_size': 20, 'unit_cost_vnd': 1000}
    row = calculate_replenishment(product, 140)
    assert (row['inventory_position'], row['reorder_point'], row['target_stock']) == (70, 70, 110)
    assert row['quantity'] == 40
    assert row['risk'] == 'Low Stock'
    product['on_hand'] += 1
    assert calculate_replenishment(product, 140)['quantity'] == 0


def test_inbound_coverage_does_not_hide_physical_stockout():
    product = {'on_hand': 0, 'on_order': 1000, 'lead_time_days': 5,
               'safety_stock': 20, 'pack_size': 20, 'unit_cost_vnd': 1000}
    row = calculate_replenishment(product, 140)
    assert row['quantity'] == 0
    assert row['risk'] == 'Critical Stock'


def test_sku_import_and_ai_use_the_same_identifier_contract(db):
    conn, _ = db
    raw = (b'sku,product_name,supplier,lead_time_days,safety_stock,pack_size,unit_cost_vnd\n'
           b'COFFEE BLEND 1,Demo,Demo supplier Alpha,5,20,10,1000\n')
    with pytest.raises(ValueError, match='Invalid SKU'):
        import_dataset(conn, raw, 'products.csv', 'products', as_of=AS_OF)


@pytest.mark.parametrize('field,value', [('on_hand', -1), ('pack_size', 0),
                                         ('reserved_qty', 1.2), ('unit_cost_vnd', True)])
def test_calculation_rejects_invalid_inputs(field, value):
    product = {'on_hand': 1, 'on_order': 0, 'lead_time_days': 5, 'safety_stock': 2,
               'pack_size': 10, 'unit_cost_vnd': 1000, field: value}
    with pytest.raises(ValueError):
        calculate_replenishment(product, 140)


def test_missing_supplier_and_missing_sales_are_persistent_exceptions(db):
    conn, _ = db
    with conn:
        conn.execute("UPDATE products SET supplier='' WHERE sku='CF-001'")
        conn.execute("DELETE FROM sales WHERE sku='RM-005'")
    rows = {r['sku']: r for r in get_replenishment_table(conn, AS_OF)}
    for sku in ('CF-001', 'RM-005'):
        assert rows[sku]['risk'] == 'Data Quality Issue'
        assert rows[sku]['quantity'] == 0
    create_drafts(conn, AS_OF)
    assert {e['sku'] for e in get_exceptions(conn)} == {'CF-001', 'RM-005'}
    create_drafts(conn, AS_OF)
    assert len(get_exceptions(conn)) == 2


def test_simulated_limit_blocks_drafts_and_recovers(db):
    conn, _ = db
    assert create_drafts(conn, AS_OF, max_value_vnd=1) == 0
    assert any(e['exception_type'] == 'POLICY_LIMIT' for e in get_exceptions(conn))
    assert create_drafts(conn, AS_OF) >= 2
    assert all(e['status'] == 'RESOLVED' for e in get_exceptions(conn))


def test_stale_draft_cannot_be_approved_but_can_be_rejected(db):
    conn, _ = db
    create_drafts(conn, AS_OF)
    proposal = get_proposals(conn)[0]
    raw = f"sku,on_hand,on_order,reserved_qty\n{proposal['sku']},1,0,0\n".encode()
    import_dataset(conn, raw, 'inventory.csv', 'inventory', as_of=AS_OF)
    with pytest.raises(ValueError, match='changed'):
        decide(conn, proposal['id'], 'APPROVED', 'reviewer', 'Review completed')
    assert not get_approval_history(conn)
    assert get_exceptions(conn)[0]['exception_type'] == 'STALE_DRAFT'
    assert conn.execute('SELECT status FROM proposals WHERE id=?', (proposal['id'],)).fetchone()[0] == 'DRAFT'
    decide(conn, proposal['id'], 'REJECTED', 'reviewer', 'Reject stale recommendation')
    assert get_exceptions(conn)[0]['status'] == 'RESOLVED'


@pytest.mark.parametrize('decision,actor,reason', [('APPROVED', '', 'Review'),
                                                 ('REJECTED', 'reviewer', ''),
                                                 ('INVALID', 'reviewer', 'Review')])
def test_invalid_review_leaves_workflow_unchanged(db, decision, actor, reason):
    conn, _ = db
    create_drafts(conn, AS_OF)
    with pytest.raises(ValueError):
        decide(conn, get_proposals(conn)[0]['id'], decision, actor, reason)
    assert not get_approval_history(conn)


def test_value_snapshot_survives_price_change(db):
    conn, _ = db
    create_drafts(conn, AS_OF)
    proposal = get_proposals(conn)[0]
    with conn:
        conn.execute('UPDATE products SET unit_cost_vnd=unit_cost_vnd+100 WHERE sku=?', (proposal['sku'],))
    assert get_proposals(conn)[0]['estimated_cost_vnd'] == proposal['estimated_cost_vnd']
    with pytest.raises(ValueError, match='changed'):
        decide(conn, proposal['id'], 'APPROVED', 'reviewer', 'Reviewed price')


def test_two_connections_cannot_generate_duplicate_drafts(db):
    conn, path = db
    def generate(_):
        other = connect(path)
        try:
            return create_drafts(other, AS_OF)
        finally:
            other.close()
    with ThreadPoolExecutor(max_workers=2) as executor:
        counts = list(executor.map(generate, range(2)))
    assert sum(counts) == len(get_proposals(conn))
    assert min(counts) == 0


def test_two_connections_cannot_approve_twice(db):
    conn, path = db
    create_drafts(conn, AS_OF)
    proposal = get_proposals(conn)[0]
    def approve(_):
        other = connect(path)
        try:
            decide(other, proposal['id'], 'APPROVED', 'reviewer', 'Concurrent review test')
            return True
        except ValueError:
            return False
        finally:
            other.close()
    with ThreadPoolExecutor(max_workers=2) as executor:
        assert sorted(executor.map(approve, range(2))) == [False, True]
    assert len(get_approval_history(conn)) == 1
    assert conn.execute('SELECT on_order FROM inventory WHERE sku=?', (proposal['sku'],)).fetchone()[0] == proposal['quantity']


def test_audit_and_approval_history_cannot_be_deleted(db):
    conn, _ = db
    create_drafts(conn, AS_OF)
    decide(conn, get_proposals(conn)[0]['id'], 'REJECTED', 'reviewer', 'Reviewed')
    for table in ('audit_logs', 'approval_history'):
        for statement in (f'DELETE FROM {table}', f"UPDATE {table} SET actor='tampered'"):
            with pytest.raises(sqlite3.IntegrityError, match='append-only'):
                with conn:
                    conn.execute(statement)


def test_full_persistent_import_approval_and_report(tmp_path):
    path = tmp_path / 'integration.db'
    conn = connect(path)
    init_db(conn)
    import_dataset(conn, b'name\nSynthetic Supplier\n', 'suppliers.csv', 'suppliers', as_of=AS_OF)
    import_dataset(conn, b'sku,product_name,supplier,lead_time_days,safety_stock,pack_size,unit_cost_vnd\n'
                   b'DEMO,Demo Material,Synthetic Supplier,5,20,20,1000\n',
                   'products.csv', 'products', as_of=AS_OF)
    import_dataset(conn, b'sku,on_hand,on_order,reserved_qty\nDEMO,55,0,5\n',
                   'inventory.csv', 'inventory', as_of=AS_OF)
    workbook = openpyxl.Workbook()
    workbook.active.append(['sale_date', 'sku', 'units'])
    workbook.active.append([AS_OF, 'DEMO', 140])
    output = io.BytesIO()
    workbook.save(output)
    assert import_dataset(conn, output.getvalue(), 'sales.xlsx', 'sales', as_of=AS_OF) == 1
    row = get_replenishment_table(conn, AS_OF)[0]
    assert row['inventory_position'] == 50
    assert row['quantity'] == 60
    assert create_drafts(conn, AS_OF) == 1
    proposal = get_proposals(conn)[0]
    decide(conn, proposal['id'], 'APPROVED', 'demo_manager', 'Confirmed synthetic shortage')
    assert create_drafts(conn, AS_OF) == 0
    conn.close()
    reopened = connect(path)
    try:
        report = build_report(reopened, AS_OF)
        assert report['proposals'][0]['status'] == 'APPROVED'
        assert report['inventory'][0]['on_order'] == 60
        event = next(a for a in report['audit_logs'] if a['action'] == 'APPROVED')
        assert json.loads(event['before_json'])['on_order'] == 0
        assert json.loads(event['after_json'])['on_order'] == 60
        assert 'APPROVED' in make_csv(report['proposals'], list(report['proposals'][0])).decode('utf-8-sig')
        exported = openpyxl.load_workbook(io.BytesIO(make_excel(report)), data_only=True)
        assert exported['approval_history'].max_row == 2
        assert exported['inventory'].max_row == 2
    finally:
        reopened.close()
